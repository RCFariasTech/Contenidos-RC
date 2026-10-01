"""Exportación del mes aprobado a PowerPoint (PLAN.md, sección 9).

Portada + una slide 16:9 por pieza, con la identidad de config/marca.json y las
medidas de config/instagram_specs.json. Todo en Arial (viene con Office: no hay
que incrustar fuentes y el ajuste de texto es predecible).
"""

import io
import math
from datetime import date

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from rc.config import RAIZ, ajustes, marca, specs_instagram
from rc.esquema import CAMPOS_TEXTO

ANCHO, ALTO = Inches(13.333), Inches(7.5)
MARGEN = Inches(0.5)
FUENTE = "Arial"
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto",
         "septiembre", "octubre", "noviembre", "diciembre"]
DIAS = ["lunes", "martes", "miércoles", "jueves", "viernes", "sábado", "domingo"]
ROLES = {
    "Carrusel": ["Gancho", "Contenido", "Contenido", "Contenido", "Cierre"],
    "Reel": ["Gancho · 0-5 s", "Desarrollo · 5-12 s", "Desarrollo · 12-20 s", "Desarrollo · 20-25 s", "CTA · 25-30 s"],
}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _colores() -> dict:
    p = marca()["paleta"]
    return {**{k: _rgb(v) for k, v in p.items()}, "tinte": _rgb("#EEF7FC"), "suave": _rgb("#4A5A78")}


def fecha_larga(iso: str | None) -> str:
    if not iso:
        return "sin fecha"
    d = date.fromisoformat(iso[:10])
    return f"{DIAS[d.weekday()]} {d.day} de {MESES[d.month - 1]}"


def nombre_mes(mes_iso: str) -> str:
    d = date.fromisoformat(mes_iso[:10])
    return f"{MESES[d.month - 1].capitalize()} {d.year}"


# ---------- ajuste de texto ----------

def tamano_que_cabe(parrafos: list[str], ancho_in: float, alto_in: float, base: float, minimo: float = 9) -> float:
    """Mayor tamaño (pt) con el que los párrafos caben en la caja. Estimación conservadora para Arial."""
    tamano = base
    while tamano > minimo:
        chars_linea = max(1, int(ancho_in * 72 / (tamano * 0.53)))
        lineas = sum(max(1, math.ceil(len(p) / chars_linea)) for p in parrafos)
        if lineas * tamano * 1.25 / 72 <= alto_in:
            return tamano
        tamano -= 0.5
    return minimo


# ---------- primitivas ----------

def _caja_texto(slide, x, y, w, h, ancla=MSO_ANCHOR.TOP):
    caja = slide.shapes.add_textbox(x, y, w, h)
    tf = caja.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = ancla
    return tf


def _parrafo(tf, texto, tamano, color, negrita=False, primero=False, espacio_despues=0, alinear=None,
             enlace=None):
    p = tf.paragraphs[0] if primero else tf.add_paragraph()
    if alinear is not None:
        p.alignment = alinear
    p.space_after = Pt(espacio_despues)
    run = p.add_run()
    run.text = texto
    f = run.font
    f.name, f.size, f.bold = FUENTE, Pt(tamano), negrita
    f.color.rgb = color
    if enlace:
        run.hyperlink.address = enlace
    return p


def _forma(slide, tipo, x, y, w, h, relleno=None, borde=None, radio=None):
    s = slide.shapes.add_shape(tipo, x, y, w, h)
    if relleno is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = relleno
    if borde is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = borde
        s.line.width = Pt(1.25)
    s.shadow.inherit = False
    if radio is not None and tipo == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radio
    return s


def _pastilla(slide, x, y, texto, relleno, color_texto, alto=Inches(0.36), tamano=11, borde=None, ancho=None):
    factor = 0.78 if texto.isupper() else 0.6
    ancho = ancho or Emu(int(Inches(0.36) + Pt(tamano) * factor * len(texto)))
    s = _forma(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, ancho, alto, relleno, borde, radio=0.5)
    tf = s.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _parrafo(tf, texto, tamano, color_texto, negrita=True, primero=True, alinear=PP_ALIGN.CENTER)
    return ancho


# ---------- slides ----------

def _portada(prs, mes_iso: str, piezas: list[dict], aprobado_en: str | None):
    c = _colores()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    fondo = s.background.fill
    fondo.gradient()
    fondo.gradient_angle = 90
    inicio, fin = marca()["gradientes_verticales"]["noche"]
    fondo.gradient_stops[0].color.rgb = _rgb(inicio)
    fondo.gradient_stops[1].color.rgb = _rgb(fin)

    s.shapes.add_picture(str(RAIZ / marca()["logos"]["fondo_oscuro"]), MARGEN, MARGEN, height=Inches(0.6))
    tf = _caja_texto(s, MARGEN, Inches(1.55), Inches(7.6), Inches(1.6))
    _parrafo(tf, "Contenido Instagram", 40, c["blanco"], negrita=True, primero=True)
    _parrafo(tf, nombre_mes(mes_iso), 28, c["celeste"], negrita=True, espacio_despues=6)
    detalle = f"Aprobado el {fecha_larga(aprobado_en)}" if aprobado_en else "Plan mensual"
    _parrafo(tf, f"{detalle} · {len(piezas)} piezas", 14, c["blanco"])

    # Resumen: una tarjeta por pieza
    y = Inches(3.75)
    alto, gap = Inches(0.66), Inches(0.14)
    for p in piezas:
        tarjeta = _forma(s, MSO_SHAPE.ROUNDED_RECTANGLE, MARGEN, y, ANCHO - 2 * MARGEN, alto,
                         relleno=c["blanco"], radio=0.18)
        tarjeta.fill.transparency = 0  # sólido: contraste máximo para el texto azul
        cx = MARGEN + Inches(0.2)
        cy = y + (alto - Inches(0.36)) // 2
        color_formato = c["azul_medio"] if p["formato"] == "Carrusel" else c["azul_marino"]
        ancho = _pastilla(s, cx, cy, p["formato"].upper(), color_formato, c["blanco"], ancho=Inches(1.25))
        tf = _caja_texto(s, cx + ancho + Inches(0.25), y, ANCHO - 2 * MARGEN - ancho - Inches(0.7), alto,
                         ancla=MSO_ANCHOR.MIDDLE)
        tema = (p.get("contenido") or {}).get("tema_especifico", "")
        _parrafo(tf, f"Semana {p['semana']} · {fecha_larga(p.get('fecha_publicacion'))}", 11, c["suave"],
                 negrita=True, primero=True)
        _parrafo(tf, tema, tamano_que_cabe([tema], 10.2, 0.3, 15, 11), c["azul_marino"], negrita=True)
        y += alto + gap
    return s


def _alto_estimado(items, ancho_in: float, escala: float) -> float:
    """Alto (pulgadas) de una lista de párrafos (texto, tamaño, …, espacio_despues) a una escala dada."""
    total = 0.0
    for texto, tamano, *_resto, espacio in items:
        t = tamano * escala
        chars_linea = max(1, int(ancho_in * 72 / (t * 0.55)))
        lineas = max(1, math.ceil(len(texto) / chars_linea))
        total += (lineas * t * 1.22 + espacio * escala) / 72
    return total


def _bloque(slide, x, y, w, h, items, minimo=0.62):
    """Escribe párrafos (texto, tamaño, color, negrita, enlace, espacio_despues) escalándolos para que quepan."""
    ancho_in, alto_in = w / 914400, h / 914400
    escala = 1.0
    while escala > minimo and _alto_estimado(items, ancho_in, escala) > alto_in:
        escala -= 0.03
    tf = _caja_texto(slide, x, y, w, h)
    for i, (texto, tamano, color, negrita, enlace, espacio) in enumerate(items):
        _parrafo(tf, texto, round(tamano * escala * 2) / 2, color, negrita=negrita, primero=(i == 0),
                 espacio_despues=espacio * escala, enlace=enlace)
    return escala


def _recortar(texto: str, maximo: int) -> str:
    texto = (texto or "").strip()
    return texto if len(texto) <= maximo else texto[: maximo - 1].rstrip() + "…"


def _ficha(slide, x, y, w, h, pieza, c):
    specs = specs_instagram()[pieza["formato"]]
    contenido = pieza["contenido"]
    inv = contenido.get("investigacion") or {}
    _forma(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, relleno=c["tinte"], radio=0.05)
    pad = Inches(0.22)
    url = inv.get("fuente_url", "")
    etiqueta = lambda t: (t.upper(), 10, c["azul_medio"], True, None, 3)  # noqa: E731
    items = [
        etiqueta("Ficha técnica"),
        (f"{pieza['formato']} · {specs['tamano_px']} px", 14, c["azul_marino"], True, None, 0),
        (f"Relación {specs['relacion']}", 11, c["azul_marino"], False, None, 3),
    ]
    for nota in filter(None, [specs.get("notas"), specs.get("portada") and f"Portada: {specs['portada']}",
                              specs.get("zona_segura") and f"Zona segura: {specs['zona_segura']}"]):
        items.append((nota, 9.5, c["suave"], False, None, 3))
    items += [
        (f"{pieza['tipo']} · {pieza['pilar']}", 10.5, c["azul_marino"], False, None, 10),
        etiqueta("Hashtags"),
        ("  ".join(contenido.get("hashtags") or []) or "Sin hashtags (ninguno aportaba alcance)", 12, c["azul_marino"], True, None, 10),
        etiqueta("Fuente"),
        (_recortar(inv.get("tendencia", ""), 220), 10.5, c["azul_marino"], True, None, 3),
        (_recortar(inv.get("fuente_titulo") or url, 120), 10, c["azul_medio"], False,
         url if url.startswith("https://") else None, 2),
        (f"Fecha: {_recortar(inv.get('fuente_fecha') or 'sin fecha visible', 60)}", 9.5, c["suave"], False, None, 0),
    ]
    _bloque(slide, x + pad, y + pad, w - 2 * pad, h - 2 * pad, items)


def _textos(slide, x, y, w, h, pieza, c):
    """Filas numeradas de alto variable según el largo de cada texto."""
    contenido = pieza["contenido"]
    campos = CAMPOS_TEXTO[pieza["formato"]]
    titulo = "Guion por escena · texto en pantalla" if pieza["formato"] == "Reel" else "Textos por tarjeta"
    _bloque(slide, x, y, w, Inches(0.3), [(titulo.upper(), 10, c["azul_medio"], True, None, 0)])

    circulo, sangria, gap = Inches(0.42), Inches(0.6), Inches(0.12)
    ancho_texto = w - sangria
    textos = [contenido.get(cp, "") for cp in campos]
    disponible = (h - Inches(0.4) - gap * (len(campos) - 1)) / 914400

    def alto_fila(texto, tamano):
        lineas = max(1, math.ceil(len(texto) / max(1, int(ancho_texto / 914400 * 72 / (tamano * 0.5)))))
        return max(circulo / 914400, (9 * 1.22 + 2 + lineas * tamano * 1.22) / 72)

    tamano = 14.0
    while tamano > 9 and sum(alto_fila(t, tamano) for t in textos) > disponible:
        tamano -= 0.5

    fila_y = y + Inches(0.4)
    for i, texto in enumerate(textos):
        alto = Inches(alto_fila(texto, tamano))
        num = _forma(slide, MSO_SHAPE.OVAL, x, fila_y, circulo, circulo, relleno=c["azul_marino"])
        ntf = num.text_frame
        ntf.margin_left = ntf.margin_right = ntf.margin_top = ntf.margin_bottom = 0
        ntf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _parrafo(ntf, str(i + 1), 13, c["blanco"], negrita=True, primero=True, alinear=PP_ALIGN.CENTER)
        tf = _caja_texto(slide, x + sangria, fila_y, ancho_texto, alto)
        etiqueta = f"{'Escena' if pieza['formato'] == 'Reel' else 'Tarjeta'} {i + 1} · {ROLES[pieza['formato']][i]}"
        _parrafo(tf, etiqueta.upper(), 9, c["suave"], negrita=True, primero=True, espacio_despues=2)
        _parrafo(tf, texto, tamano, c["azul_marino"], negrita=(i == 0))
        fila_y += alto + gap


def _caption(slide, x, y, w, h, contenido, c):
    cta = ajustes()["cta"]
    cuerpo = (contenido.get("caption") or "").strip()
    if cuerpo.endswith(cta):  # el CTA va en su propia línea, aunque el modelo lo pegue al último párrafo
        cuerpo = cuerpo[: -len(cta)].strip()
    items = [("CAPTION", 10, c["azul_medio"], True, None, 4)]
    items += [(t.strip(), 11.5, c["azul_marino"], False, None, 5) for t in cuerpo.split("\n") if t.strip()]
    items.append((cta, 11.5, c["azul_marino"], True, None, 0))
    _bloque(slide, x, y, w, h, items, minimo=0.7)


def _slide_pieza(prs, pieza: dict):
    c = _colores()
    s = prs.slides.add_slide(prs.slide_layouts[6])
    contenido = pieza["contenido"]
    ancho_util = ANCHO - 2 * MARGEN

    s.shapes.add_picture(str(RAIZ / marca()["logos"]["fondo_claro"]), ANCHO - MARGEN - Inches(1.75),
                         MARGEN, height=Inches(0.5))
    tf = _caja_texto(s, MARGEN, MARGEN, ancho_util - Inches(2.1), Inches(0.3))
    _parrafo(tf, f"SEMANA {pieza['semana']} · PUBLICACIÓN: {fecha_larga(pieza.get('fecha_publicacion')).upper()}",
             12, c["azul_medio"], negrita=True, primero=True)

    tema = contenido.get("tema_especifico", "")
    ancho_titulo = ancho_util - Inches(2.1)
    tf = _caja_texto(s, MARGEN, MARGEN + Inches(0.38), ancho_titulo, Inches(0.95))
    _parrafo(tf, tema, tamano_que_cabe([tema], ancho_titulo / 914400, 0.95, 28, 18), c["azul_marino"],
             negrita=True, primero=True)

    specs = specs_instagram()[pieza["formato"]]
    y_pastillas = MARGEN + Inches(1.42)
    color_formato = c["azul_medio"] if pieza["formato"] == "Carrusel" else c["azul_marino"]
    ancho = _pastilla(s, MARGEN, y_pastillas, pieza["formato"].upper(), color_formato, c["blanco"])
    ancho2 = _pastilla(s, MARGEN + ancho + Inches(0.15), y_pastillas,
                       f"{specs['tamano_px']} px · {specs['relacion']}", c["blanco"], c["azul_marino"],
                       borde=c["azul_marino"])
    _pastilla(s, MARGEN + ancho + ancho2 + Inches(0.3), y_pastillas, pieza["tipo"], c["blanco"], c["azul_marino"],
              borde=c["coral"])

    top = MARGEN + Inches(2.05)
    alto_cuerpo = ALTO - MARGEN - top
    separacion = Inches(0.35)
    ancho_ficha, ancho_caption = Inches(3.45), Inches(3.75)
    ancho_textos = ancho_util - ancho_ficha - ancho_caption - 2 * separacion
    _ficha(s, MARGEN, top, ancho_ficha, alto_cuerpo, pieza, c)
    x_textos = MARGEN + ancho_ficha + separacion
    _textos(s, x_textos, top, ancho_textos, alto_cuerpo, pieza, c)
    _caption(s, x_textos + ancho_textos + separacion, top, ancho_caption, alto_cuerpo, contenido, c)

    inv = contenido.get("investigacion") or {}
    s.notes_slide.notes_text_frame.text = (
        f"Tendencia: {inv.get('tendencia', '')}\n\n{inv.get('resumen', '')}\n\n"
        f"Estrategia clave: {inv.get('estrategia_clave', '')}\n\nFuente: {inv.get('fuente_url', '')}")
    return s


def generar_pptx(mes_iso: str, piezas: list[dict], aprobado_en: str | None = None) -> bytes:
    """piezas: filas de la tabla piezas (con contenido), en cualquier orden."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = ANCHO, ALTO
    prs.core_properties.title = f"Contenido Instagram RC Farias · {nombre_mes(mes_iso)}"
    prs.core_properties.author = "RC Farias"
    ordenadas = sorted(piezas, key=lambda p: p["semana"])
    _portada(prs, mes_iso, ordenadas, aprobado_en)
    for p in ordenadas:
        _slide_pieza(prs, p)
    salida = io.BytesIO()
    prs.save(salida)
    return salida.getvalue()


def nombre_archivo(mes_iso: str) -> str:
    return f"RC_Farias_Instagram_{mes_iso[:7]}.pptx"
