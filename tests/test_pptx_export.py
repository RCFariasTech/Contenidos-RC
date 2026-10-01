"""El PowerPoint se genera desde datos reales del mes y contiene lo que el equipo necesita."""

import io
import json
import unittest
from pathlib import Path

from pptx import Presentation

from rc.config import ajustes
from rc.esquema import CAMPOS_TEXTO
from rc.pptx_export import fecha_larga, generar_pptx, nombre_archivo, tamano_que_cabe

FIXTURE = Path(__file__).parent / "fixtures" / "mes_noviembre.json"


def textos_de(slide) -> str:
    return "\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)


class TestPptx(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.datos = json.loads(FIXTURE.read_text(encoding="utf-8"))
        contenido = generar_pptx(cls.datos["mes"], cls.datos["piezas"], cls.datos["aprobado_en"])
        cls.prs = Presentation(io.BytesIO(contenido))

    def test_portada_y_una_slide_por_pieza(self):
        self.assertEqual(len(self.prs.slides), 1 + len(self.datos["piezas"]))
        portada = textos_de(self.prs.slides[0])
        self.assertIn("Noviembre 2026", portada)
        for p in self.datos["piezas"]:
            self.assertIn(p["contenido"]["tema_especifico"], portada)

    def test_cada_slide_tiene_formato_medidas_textos_caption_y_hashtags(self):
        medidas = {"Carrusel": "1080 × 1440", "Reel": "1080 × 1920"}
        for slide, pieza in zip(list(self.prs.slides)[1:], sorted(self.datos["piezas"], key=lambda p: p["semana"])):
            texto = textos_de(slide)
            c = pieza["contenido"]
            self.assertIn(medidas[pieza["formato"]], texto)
            self.assertIn(fecha_larga(pieza["fecha_publicacion"]).upper(), texto)
            for campo in CAMPOS_TEXTO[pieza["formato"]]:
                self.assertIn(c[campo], texto)
            for h in c["hashtags"]:
                self.assertIn(h, texto)
            self.assertIn(ajustes()["cta"], texto)
            self.assertIn(c["caption"][:60], texto)
            self.assertIn(c["investigacion"]["resumen"], slide.notes_slide.notes_text_frame.text)

    def test_todo_dentro_de_la_slide(self):
        ancho, alto = self.prs.slide_width, self.prs.slide_height
        for i, slide in enumerate(self.prs.slides):
            for s in slide.shapes:
                self.assertLessEqual(s.left + s.width, ancho, f"slide {i + 1}: {s.name} se sale por la derecha")
                self.assertLessEqual(s.top + s.height, alto, f"slide {i + 1}: {s.name} se sale por abajo")

    def test_utilidades(self):
        self.assertEqual(fecha_larga("2026-11-03"), "martes 3 de noviembre")
        self.assertEqual(nombre_archivo("2026-11-01"), "RC_Farias_Instagram_2026-11.pptx")
        self.assertLess(tamano_que_cabe(["x" * 2000], 3, 1, 14), 14)
        self.assertEqual(tamano_que_cabe(["corto"], 3, 1, 14), 14)


if __name__ == "__main__":
    unittest.main()
